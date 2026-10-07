from __future__ import annotations

import argparse
import contextlib
import importlib
import json
import sys
import tempfile
from pathlib import Path

import numpy as np
import tifffile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

swc = importlib.import_module("02_axon_metrics")
gray = importlib.import_module("03_gray_metrics")

_EXPECTED_LABEL_SHAPE = (1214, 323, 235)
_Z_GRID = np.arange(1140, 5992 + 1, 4, dtype=np.int32)
_Y_GRID = np.arange(20, 1310, 4, dtype=np.int32)
_X_GRID = np.arange(0, 939, 4, dtype=np.int32)
_DISPLAY_SCALE_XYZ = np.asarray([7.1, 6.0, 7.1], dtype=np.float64)
_SEGMENTS = [
    ("Sacral", 530.0, 1300.0, "[530, 1300)"),
    ("Lumbar", 1300.0, 2750.0, "[1300, 2750)"),
    ("Thoracic", 2750.0, 4500.0, "[2750, 4500)"),
    ("Cervical", 4500.0, 6000.0, "[4500, 6000]"),
]


def _load_label_tiff(path: Path) -> np.ndarray:
    if not path.is_file():
        raise FileNotFoundError(path)
    labels = np.asarray(tifffile.imread(path))
    if labels.shape != _EXPECTED_LABEL_SHAPE:
        raise ValueError(f"Expected label TIFF shape {_EXPECTED_LABEL_SHAPE}, got {labels.shape}")
    if not np.issubdtype(labels.dtype, np.integer):
        raise TypeError(f"Label TIFF must contain integer values, got {labels.dtype}")
    codes = sorted(int(value) for value in np.unique(labels))
    if any(value not in (0, 1, 2, 3) for value in codes):
        raise ValueError(f"Allowed label values are 0,1,2,3; found {codes}")
    if not all(value in codes for value in (1, 2, 3)):
        raise ValueError(f"All three gray-matter regions must be present; found {codes}")
    return labels.astype(np.uint8, copy=False)


@contextlib.contextmanager
def _argv(program: str, arguments: list[str]):
    previous = sys.argv[:]
    sys.argv = [program, *arguments]
    try:
        yield
    finally:
        sys.argv = previous


def _configure(modules: list[object], swc_dir: Path, work: Path, neurons: list[tuple[str, str]]) -> None:
    for module in modules:
        if hasattr(module, "SWC_DIR"):
            module.SWC_DIR = swc_dir
        if hasattr(module, "OUTPUT_DIR"):
            module.OUTPUT_DIR = work
        if hasattr(module, "OUTPUT_PATH"):
            module.OUTPUT_PATH = work / Path(module.OUTPUT_PATH).name
        if hasattr(module, "TARGETS"):
            module.TARGETS = list(neurons)
        if hasattr(module, "REGIONS"):
            module.REGIONS = list(_SEGMENTS)
        if hasattr(module, "LONGITUDINAL_SCALE"):
            module.LONGITUDINAL_SCALE = 7.1
        if hasattr(module, "DISPLAY_SCALE_XYZ"):
            module.DISPLAY_SCALE_XYZ = _DISPLAY_SCALE_XYZ.copy()


def _scalar(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _write_records(workbook: Workbook, title: str, records: list[dict]) -> None:
    sheet = workbook.create_sheet(title=title[:31])
    if not records:
        sheet.append(["No records"])
        return
    headers: list[str] = []
    for record in records:
        for key in record:
            if key not in headers:
                headers.append(key)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for record in records:
        sheet.append([_scalar(record.get(header)) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(record.get(header, "")) for record in records[:200]]
        sheet.column_dimensions[get_column_letter(index)].width = min(42, max(10, max(map(len, values)) + 2))


def _write_workbooks(work: Path, output_dir: Path, label_tiff: Path) -> tuple[Path, Path]:
    segment = json.loads((work / "axon_segment_analysis.json").read_text(encoding="utf-8"))
    extended = json.loads((work / "axon_extended_analysis.json").read_text(encoding="utf-8"))
    gray = json.loads((work / "gray_region_analysis_final.json").read_text(encoding="utf-8"))
    crossed = json.loads((work / "gray_region_by_spinal_segment_final.json").read_text(encoding="utf-8"))

    spinal_path = output_dir / "four_neurons_whole_spinal_cord_and_four_segment_axon_statistics.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_records(workbook, "Overall metrics", extended["overview_rows"])
    _write_records(workbook, "Four-segment statistics", extended["region_rows"])
    _write_records(workbook, "Longitudinal bins", extended["bin_rows"])
    _write_records(workbook, "Axon terminals", extended["terminal_rows"])
    _write_records(workbook, "Branch points", extended["branch_point_rows"])
    _write_records(workbook, "SWC and soma information", segment["diagnostics"])
    _write_records(workbook, "Parameters", [
        {"source": "four_segments", "parameter": key, "value": _scalar(value)}
        for key, value in segment.get("parameters", {}).items()
    ] + [
        {"source": "extended", "parameter": key, "value": _scalar(value)}
        for key, value in extended.get("parameters", {}).items()
    ])
    _write_records(workbook, "Quality control", extended["diagnostics"])
    workbook.save(spinal_path)

    gray_path = output_dir / "four_neurons_final_gray_matter_region_axon_statistics.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_records(workbook, "Three gray-matter regions", gray["region_rows"])
    _write_records(workbook, "Three regions x four segments", crossed["cell_rows"])
    _write_records(workbook, "Three-region overview", gray["overview_rows"])
    _write_records(workbook, "Crossed overview", crossed["overview_rows"])
    _write_records(workbook, "Terminal details", crossed["terminal_rows"])
    _write_records(workbook, "Branch-point details", crossed["branch_point_rows"])
    _write_records(workbook, "Three-region volumes", gray["region_volume"])
    _write_records(workbook, "12-cell volumes", crossed["cell_volume_rows"])
    _write_records(workbook, "Parameters", [
        {"source": "private_label_tiff", "parameter": "label_tiff", "value": label_tiff.name},
    ] + [
        {"source": "three_zones", "parameter": key, "value": _scalar(value)}
        for key, value in gray.get("parameters", {}).items() if "label" not in key.lower()
    ] + [
        {"source": "three_by_four", "parameter": key, "value": _scalar(value)}
        for key, value in crossed.get("parameters", {}).items() if "label" not in key.lower()
    ])
    _write_records(workbook, "Quality control", [
        {"level": "three_zones", **row} for row in gray["diagnostics"]
    ] + [
        {"level": "three_by_four", **row} for row in crossed["diagnostics"]
    ])
    workbook.save(gray_path)
    return spinal_path, gray_path


def run_analysis(swc_dir: Path, gray_label_tiff: Path, output_dir: Path, neurons: list[tuple[str, str]]) -> tuple[Path, Path]:
    missing = [swc_dir / filename for _, filename in neurons if not (swc_dir / filename).is_file()]
    if missing:
        raise FileNotFoundError("Missing SWC files:\n" + "\n".join(map(str, missing)))
    labels = _load_label_tiff(gray_label_tiff)
    output_dir.mkdir(parents=True, exist_ok=True)
    modules = [swc, gray]

    with tempfile.TemporaryDirectory(prefix="spinal_axon_statistics_") as temporary:
        work = Path(temporary)
        label_npz = work / "private_gray_region_labels.npz"
        np.savez_compressed(label_npz, labels=labels, z=_Z_GRID, y=_Y_GRID, x=_X_GRID)
        del labels
        _configure(modules, swc_dir, work, neurons)
        swc.OUTPUT_DIR = work
        swc.OUTPUT_PATH = work / "axon_extended_analysis.json"
        swc.run_segment_analysis()
        swc.run_extended_analysis()
        gray_output = work / "gray_region_analysis_final.json"
        cross_output = work / "gray_region_by_spinal_segment_final.json"
        with _argv(gray.__name__, [
            "--labels", str(label_npz), "--output", str(gray_output),
            "--partition-description", "Private finalized dorsal/middle/ventral gray-matter labels",
        ]):
            gray.run_gray_analysis()
        with _argv(gray.__name__, [
            "--labels", str(label_npz), "--output", str(cross_output),
            "--version-name", "final", "--partition-description",
            "Private finalized dorsal/middle/ventral gray-matter labels",
            "--reference-analysis", str(gray_output),
        ]):
            gray.run_gray_segment_analysis()
        outputs = _write_workbooks(work, output_dir, gray_label_tiff)
    for path in outputs:
        print(path)
    return outputs


SWC_DIR = Path(r"PATH/TO/SWC_FILES")
GRAY_REGION_LABEL_TIFF = Path(r"PATH/TO/PRIVATE_GRAY_REGION_LABELS.tif")
OUTPUT_DIR = Path(__file__).resolve().parent / "results"

NEURONS = [
    ("strong_central_3", "strong_central_3.swc"),
    ("stong-4_central", "stong-4_central_done.swc"),
    ("tree-10_dorsal", "tree-10_dorsal_done.swc"),
    ("tree-11_dorsal", "tree-11_dorsal_done.swc"),
]

def main() -> None:
    parser = argparse.ArgumentParser(description="Quantify four SWC axons using a private three-zone gray-matter label TIFF.")
    parser.add_argument("--swc-dir", type=Path, default=SWC_DIR)
    parser.add_argument("--gray-label-tiff", type=Path, default=GRAY_REGION_LABEL_TIFF)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()
    run_analysis(
        swc_dir=args.swc_dir.expanduser().resolve(),
        gray_label_tiff=args.gray_label_tiff.expanduser().resolve(),
        output_dir=args.output_dir.expanduser().resolve(),
        neurons=NEURONS,
    )

if __name__ == "__main__":
    main()
